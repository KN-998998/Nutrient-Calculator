"""解析 3 份历史营养结果文件为统一结构, 用于与工具计算结果比对。

输出: history_results.json
  [{"month": "4", "name": "A鮮茄雞蛋肉片配飯...", "nutrients": {"energy_kcal": ..., ...}}, ...]
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

FOLDER = Path(r"F:\营养素计算工具\菜单+SOP")
OUT = Path(r"F:\营养素计算工具\data\history_results.json")

# 统一营养字段映射: (正则模式, 工具字段名)
PATTERNS = [
    (r"能量\s*([\d.]+)\s*[（(]?kcal", "energy_kcal"),
    (r"蛋白質\s*([\d.]+)\s*g", "protein_g"),
    (r"總脂肪\s*([\d.]+)\s*g", "total_fat_g"),
    (r"飽和脂肪\s*([\d.]+)\s*g", "saturated_fat_g"),
    (r"反式脂肪\s*([\d.]+)\s*g", "trans_fat_g"),
    (r"碳水化合物\s*([\d.]+)\s*g", "available_carbohydrate_g"),
    (r"糖\s*([\d.]+)\s*g", "sugars_g"),
    (r"鈉\s*([\d.]+)\s*mg", "sodium_mg"),
]


def parse_html_detail(text: str) -> dict:
    nutrients = {}
    for pattern, field in PATTERNS:
        m = re.search(pattern, text)
        if m:
            nutrients[field] = float(m.group(1))
    return nutrients


def parse_april(sheet_name: str = "Sheet1") -> list[dict]:
    wb = openpyxl.load_workbook(FOLDER / "4月菜式营养素.xlsx", read_only=True, data_only=True)
    ws = wb[sheet_name]
    results = []
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    # 4月表头: [商品名称, 能量（kcal）|值, 能量（kcal）|值(重复), 蛋白質（g）|值, ...]
    # 实际结构: col0=名称, 然后每项营养两列交替: col1=标签, col2=值, col3=标签, col4=值...
    # 数值列: col2, col4, col6, ..., col16 (8 项)
    LABEL_COLS = [
        (1, 2, "能量", "energy_kcal"),
        (3, 4, "蛋白質", "protein_g"),
        (5, 6, "總脂肪", "total_fat_g"),
        (7, 8, "飽和脂肪", "saturated_fat_g"),
        (9, 10, "反式脂肪", "trans_fat_g"),
        (11, 12, "碳水化合物", "available_carbohydrate_g"),
        (13, 14, "糖", "sugars_g"),
        (15, 16, "鈉", "sodium_mg"),
    ]
    for row in rows[1:]:
        name = str(row[0] or "").strip()
        if not name:
            continue
        nutrients = {}
        for label_col, value_col, label_kw, field in LABEL_COLS:
            if len(row) <= value_col:
                continue
            label = str(row[label_col] or "").strip()
            value = row[value_col]
            if label_kw not in label or value is None:
                continue
            try:
                nutrients[field] = float(value)
            except (TypeError, ValueError):
                continue
        results.append({"month": "4", "name": name, "nutrients": nutrients})
    return results


def parse_may() -> list[dict]:
    wb = openpyxl.load_workbook(FOLDER / "5月學生餐菜品營養素5.28.xlsx", read_only=True, data_only=True)
    ws = wb["5月菜品營養素"]
    results = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = str(row[0] or "").strip()
        detail = str(row[1] or "")
        if not name:
            continue
        results.append({
            "month": "5",
            "name": name,
            "nutrients": parse_html_detail(detail),
        })
    wb.close()
    return results


def parse_june() -> list[dict]:
    wb = openpyxl.load_workbook(FOLDER / "6月菜品營養素0702.xlsx", read_only=True, data_only=True)
    ws = wb["6月菜品營養素"]
    results = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = str(row[0] or "").strip()
        detail = str(row[4] or "")  # 商品详情列
        if not name:
            continue
        results.append({
            "month": "6",
            "name": name,
            "nutrients": parse_html_detail(detail),
        })
    wb.close()
    return results


def main() -> int:
    all_results = parse_april() + parse_may() + parse_june()
    OUT.write_text(json.dumps(all_results, ensure_ascii=False, indent=1), encoding="utf-8")

    # 统计
    total = len(all_results)
    with_nutrients = sum(1 for r in all_results if r["nutrients"])
    complete = sum(1 for r in all_results if len(r["nutrients"]) == 8)
    by_month = {}
    for r in all_results:
        by_month.setdefault(r["month"], 0)
        by_month[r["month"]] += 1
    print(f"解析 {total} 条 (4月 {by_month.get('4',0)} / 5月 {by_month.get('5',0)} / 6月 {by_month.get('6',0)})")
    print(f"有营养值: {with_nutrients}, 8项齐全: {complete}")
    print(f"输出: {OUT}")
    # 抽查
    for r in all_results[:3]:
        print(f"  [{r['month']}月] {r['name'][:30]} -> {r['nutrients']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
