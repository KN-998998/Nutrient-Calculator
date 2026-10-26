"""Build the storefront upload workbook (only columns A and E filled).

The target workbook is the 15-column storefront template (e.g.
"聖類斯上架模版-9.15.xlsx").  This script:
- reads the monthly menu (default: 9月主菜單) and selects one dish per date/meal;
- matches each menu dish to the SOP recipe by normalized name;
- calculates the eight Hong Kong nutrients per school stage;
- writes one sheet per stage using the supplied template, filling only
  column A (product name, e.g. A(1)菜名) and column E (nutrition detail);
- leaves every other column blank.

Usage:
    python -X utf8 scripts/build_store_upload.py ^
        "D:/.../2026年中小學生9月午餐餐單sop-10.10.xlsx" ^
        "D:/.../聖類斯上架模版-9.15.xlsx" ^
        "outputs/2026年9月午餐_上架模板_A-D.xlsx"
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import difflib
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from nutrition_tool.calculator import calculate_recipe
from nutrition_tool.dish_export import (
    NUTRIENT_TO_DISH_COL,
    format_store_nutrition_detail,
    write_store_upload_excel,
)
from nutrition_tool.sop import read_sop_workbook
from web_app.app import get_repo


MARKERS = re.compile(r"[#★●▲□◆※♁◎]")
NOTE_PATTERN = re.compile(r"含微量魚骨|不含蛋|\(不含蛋\)|（不含蛋）")
PAREN_PATTERN = re.compile(r"\([^)]*\)|（[^）]*）")
SPACE_PATTERN = re.compile(r"[\s<>\[\]【】_/\-+*#%&|，,。.、]+")


def clean_product_name(value: object) -> str:
    text = str(value or "").replace("\r", " ").replace("\n", " ")
    text = MARKERS.sub("", text)
    text = NOTE_PATTERN.sub("", text)
    text = re.sub(r"\s+", "", text)
    return text.strip()


def match_key(value: object) -> str:
    text = clean_product_name(value)
    text = PAREN_PATTERN.sub("", text)
    text = SPACE_PATTERN.sub("", text)
    return text.lower().replace("臺", "台").replace("鷄", "雞")


def read_menu(path: str | Path, sheet_name: str, meals: list[str]) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"菜单工作表不存在: {sheet_name}")
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        header_index = None
        for index, row in enumerate(rows[:10]):
            if any(str(cell or "").strip() == "日期" for cell in row):
                header_index = index
                break
        if header_index is None:
            raise ValueError("菜单中找不到「日期」表头")
        header = [str(cell or "").strip() for cell in rows[header_index]]
        date_col = next(i for i, cell in enumerate(header) if cell == "日期")
        meal_cols: dict[str, int] = {}
        for meal in meals:
            for index, cell in enumerate(header):
                if re.match(rf"^{re.escape(meal)}\s*款", cell):
                    meal_cols[meal] = index
                    break
            if meal not in meal_cols:
                raise ValueError(f"菜单中找不到 {meal} 款列")
        entries = []
        for row in rows[header_index + 1:]:
            date = row[date_col] if date_col < len(row) else None
            if not isinstance(date, dt.datetime):
                continue
            for meal in meals:
                index = meal_cols[meal]
                name = row[index] if index < len(row) else None
                if name in (None, ""):
                    continue
                entries.append({
                    "date": date.date(),
                    "meal": meal,
                    "name": str(name).strip(),
                })
        return entries
    finally:
        workbook.close()


def select_recipes(entries: list[dict], recipes: list, stages: list[str]) -> tuple[list[dict], list[dict]]:
    selected: list[dict] = []
    issues: list[dict] = []
    for stage in stages:
        for entry in entries:
            candidates = [
                recipe for recipe in recipes
                if recipe.school_stage == stage
                and recipe.service_date == entry["date"].isoformat()
                and recipe.meal_set == entry["meal"]
            ]
            if not candidates:
                issues.append({**entry, "stage": stage, "reason": "没有同日期/餐别的SOP菜式", "score": 0})
                continue
            target_key = match_key(entry["name"])
            scored = sorted(
                (
                    (difflib.SequenceMatcher(None, target_key, match_key(recipe.recipe_name)).ratio(), recipe)
                    for recipe in candidates
                ),
                key=lambda item: item[0],
                reverse=True,
            )
            score, recipe = scored[0]
            if score < 0.55:
                issues.append({
                    **entry, "stage": stage,
                    "reason": f"名称匹配度过低，候选={recipe.recipe_name}",
                    "score": round(score, 3),
                })
                continue
            selected.append({
                **entry,
                "stage": stage,
                "recipe": recipe,
                "score": round(score, 3),
            })
    return selected, issues


def build_rows(selected: list[dict], results: dict, digits: int = 0) -> tuple[list[dict], list[dict]]:
    """Build storefront rows; results are keyed by recipe object identity to
    avoid collisions between standard and school-specific D recipes with the
    same display name, date and meal set."""
    rows: list[dict] = []
    skipped: list[dict] = []
    for item in selected:
        stage = item["stage"]
        recipe = item["recipe"]
        result = results[id(recipe)]
        product_name = f"{item['meal']}({item['date'].day}){clean_product_name(item['name'])}"
        if result.status == "incomplete":
            detail = ""
            skipped.append({**item, "reason": "SOP 食材不完整"})
        else:
            detail_row = {
                column: result.nutrients_per_serving[field]
                for field, column in NUTRIENT_TO_DISH_COL.items()
            }
            detail = format_store_nutrition_detail(detail_row, digits=digits)
        rows.append({
            "學段": stage,
            "商品名称": product_name,
            "商品详情": detail,
            "常規菜單": result.status,
            "date": item["date"].isoformat(),
            "meal": item["meal"],
            "menu_name": item["name"],
            "sop_name": recipe.recipe_name,
            "score": item["score"],
        })
    meal_order = {"A": 0, "B": 1, "C": 2, "D": 3, "E": 4}
    rows.sort(key=lambda row: (meal_order.get(row["meal"], 9), row["學段"]))
    rows.sort(key=lambda row: row["date"], reverse=True)
    return rows, skipped


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="生成上架模板：仅填 A 列商品名称、E 列营养详情")
    parser.add_argument("sop", type=Path)
    parser.add_argument("template", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--menu-sheet", default="9月主菜單")
    parser.add_argument("--meals", default="A,B,C,D")
    parser.add_argument("--stages", default="小學,中學")
    parser.add_argument("--decimals", type=int, default=0, help="E列营养值小数位，默认0（与上架模板一致）")
    parser.add_argument("--audit-csv", type=Path, help="可选：输出菜品匹配审计CSV")
    args = parser.parse_args(argv)

    meals = [item.strip().upper() for item in args.meals.split(",") if item.strip()]
    stages = [item.strip() for item in args.stages.split(",") if item.strip()]

    recipes, parse_errors = read_sop_workbook(args.sop)
    repository = get_repo("hong_kong")
    calculated = [calculate_recipe(recipe, repository) for recipe in recipes]
    results = {
        id(recipe): result
        for recipe, result in zip(recipes, calculated)
    }
    entries = read_menu(args.sop, args.menu_sheet, meals)
    selected, issues = select_recipes(entries, recipes, stages)
    rows, skipped = build_rows(selected, results, digits=args.decimals)
    write_store_upload_excel(rows, str(args.output), str(args.template))

    if args.audit_csv:
        with args.audit_csv.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=[
                "學段", "date", "meal", "商品名称", "menu_name", "sop_name", "score",
                "商品详情", "常規菜單",
            ])
            writer.writeheader()
            writer.writerows(rows)

    summary = {
        "menu_entries": len(entries),
        "selected": len(selected),
        "written": sum(1 for row in rows if row["商品详情"]),
        "skipped": len(skipped) + len(issues),
        "issues": issues,
        "parse_errors": len(parse_errors),
        "output": str(args.output),
    }
    print(summary)
    return 0 if not issues and not skipped else 1


if __name__ == "__main__":
    raise SystemExit(main())
