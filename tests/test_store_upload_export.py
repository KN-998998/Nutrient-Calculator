import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from nutrition_tool.dish_export import (
    STORE_UPLOAD_HEADERS,
    format_store_nutrition_detail,
    write_store_upload_excel,
)


class StoreNutritionFormatTests(unittest.TestCase):
    def test_zero_and_trailing_zero_values_are_preserved(self):
        detail = format_store_nutrition_detail({
            "能量(kcal)": 860.0,
            "蛋白質(g)": 38.273,
            "碳水化合物(g)": 90.473,
            "總脂肪(g)": 38.5,
            "糖(g)": 0.0,
            "飽和脂肪(g)": 8.09,
            "反式脂肪(g)": 0.0,
            "鈉(mg)": 2698.68,
        }, digits=0)

        self.assertIn("能量860(Kcal)", detail)
        self.assertIn("碳水化合物90(g)", detail)
        self.assertIn("糖0(g)", detail)
        self.assertIn("反式脂肪0(g)", detail)
        self.assertIn("鈉2699(mg)", detail)
        self.assertNotEqual(detail, "")

    def test_decimal_mode_keeps_integer_trailing_zero(self):
        detail = format_store_nutrition_detail({
            "能量(kcal)": 90.0,
            "蛋白質(g)": 10.0,
            "碳水化合物(g)": 0.0,
            "總脂肪(g)": 0.0,
            "糖(g)": 0.0,
            "飽和脂肪(g)": 0.0,
            "反式脂肪(g)": 0.0,
            "鈉(mg)": 0.0,
        }, digits=2)

        self.assertIn("能量90(Kcal)", detail)
        self.assertIn("蛋白質10(g)", detail)
        self.assertIn("碳水化合物0(g)", detail)


def _template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "upload"
    sheet.append(list(STORE_UPLOAD_HEADERS))
    sheet.append(["A(30)示例菜", "分类", "图片", "规格", "能量1(Kcal)", "0", "0", "-1", "0", "0", "1", "1", "0", "", ""])
    workbook.save(path)
    workbook.close()


class StoreUploadWorkbookTests(unittest.TestCase):
    def test_only_a_and_e_are_written_and_other_columns_are_blank(self):
        with tempfile.TemporaryDirectory() as folder:
            template = Path(folder) / "template.xlsx"
            output = Path(folder) / "upload.xlsx"
            _template(template)
            write_store_upload_excel(
                [
                    {
                        "學段": "小學",
                        "商品名称": "A(1)馬蹄蒸肉餅配飯(菜心苗)",
                        "商品详情": "能量666(Kcal) 蛋白質36(g) 碳水化合物92(g) 脂肪17(g) 糖2(g) 飽和脂肪5(g) 反式脂肪0(g) 鈉2271(mg)",
                    },
                    {
                        "學段": "中學",
                        "商品名称": "A(1)馬蹄蒸肉餅配飯(菜心苗)",
                        "商品详情": "能量715(Kcal) 蛋白質42(g) 碳水化合物92(g) 脂肪20(g) 糖2(g) 飽和脂肪6(g) 反式脂肪0(g) 鈉2286(mg)",
                    },
                ],
                str(output),
                str(template),
            )
            workbook = load_workbook(output, read_only=False, data_only=True)
            self.assertEqual(workbook.sheetnames, ["小学标准", "中学标准"])
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                self.assertEqual(sheet.max_row, 2)
                self.assertEqual(sheet.max_column, 15)
                self.assertTrue(sheet["A2"].value.startswith("A(1)"))
                self.assertTrue(sheet["E2"].value.startswith("能量"))
                for column in (2, 3, 4, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15):
                    self.assertIn(sheet.cell(2, column).value, (None, ""))
            workbook.close()


if __name__ == "__main__":
    unittest.main()
