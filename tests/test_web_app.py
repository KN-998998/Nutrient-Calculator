import unittest
import tempfile
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from web_app.app import app
from nutrition_tool.models import NutrientProfile
from nutrition_tool.repository import NutrientRepository


class WebAppHealthTests(unittest.TestCase):
    def test_healthz_returns_ok(self):
        response = app.test_client().get("/healthz")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json(), {"status": "ok"})

    def test_process_exports_primary_and_secondary_results(self):
        source = BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "9月A餐配料sop"
        sheet.append([
            "編號/圖片", "套餐", "菜式", "食材/系統名稱", "安士（小學）",
            "人均餐盒實際下單用量（g）（小學）", "安士（中學）",
            "人均餐盒實際下單用量（g）（中學）", "編碼",
        ])
        sheet.append([46266, "雞肉飯", "主菜", "雞肉", 1, 100, 2, 200, "CHICKEN"])
        workbook.save(source)
        workbook.close()
        source.seek(0)

        repository = NutrientRepository([
            NutrientProfile(
                ingredient_code="CHICKEN", ingredient_name="雞肉", energy_kcal=100,
                protein_g=20, total_fat_g=1, saturated_fat_g=0.5,
                available_carbohydrate_g=2, sugars_g=1, sodium_mg=50,
            )
        ])
        with tempfile.TemporaryDirectory(prefix="nutri_out_") as folder, \
                patch("web_app.app.tempfile.mkdtemp", return_value=folder), \
                patch("web_app.app._cleanup_old_tmpdirs"), \
                patch("web_app.app.get_history_repo", return_value=SimpleNamespace(find_confirmed=lambda recipe: None)), \
                patch("web_app.app.get_repo", return_value=repository):
            client = app.test_client()
            response = client.post(
                "/api/process",
                data={"sop_files": (source, "sample.xlsx")},
                content_type="multipart/form-data",
            )

            self.assertEqual(response.status_code, 200)
            payload = response.get_json()
            self.assertEqual(payload["incomplete"], 0)
            download = client.get(payload["download"])
            self.assertEqual(download.status_code, 200)
            result_workbook = load_workbook(BytesIO(download.data), read_only=True, data_only=True)
            self.assertEqual(result_workbook.sheetnames, ["小学", "中学"])
            self.assertEqual(result_workbook["小学"].max_column, 9)
            self.assertEqual(result_workbook["小学"]["A2"].value, "雞肉飯")
            self.assertEqual(result_workbook["小学"]["B2"].value, 100)
            self.assertEqual(result_workbook["中学"]["B2"].value, 200)
            result_workbook.close()
            download.close()

    def test_process_reports_specific_missing_items(self):
        source = BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "9月A餐配料sop"
        sheet.append([
            "編號/圖片", "套餐", "菜式", "食材/系統名稱", "安士（小學）",
            "人均餐盒實際下單用量（g）（小學）", "安士（中學）",
            "人均餐盒實際下單用量（g）（中学）", "編碼",
        ])
        sheet.append([46266, "待补数据饭", "主菜", "未知酱料", 1, 10, 1, 20, "SAUCE"])
        workbook.save(source)
        workbook.close()
        source.seek(0)

        repository = NutrientRepository([])
        with tempfile.TemporaryDirectory(prefix="nutri_out_") as folder, \
                patch("web_app.app.tempfile.mkdtemp", return_value=folder), \
                patch("web_app.app._cleanup_old_tmpdirs"), \
                patch("web_app.app.get_history_repo", return_value=SimpleNamespace(find_confirmed=lambda recipe: None)), \
                patch("web_app.app.get_repo", return_value=repository):
            with app.test_client() as client:
                response = client.post(
                    "/api/process",
                    data={"sop_files": (source, "missing.xlsx")},
                    content_type="multipart/form-data",
                )
                self.assertEqual(response.status_code, 200)
                payload = response.get_json()
                self.assertEqual(payload["incomplete"], 2)
                download = client.get(payload["download"])
                self.assertEqual(download.status_code, 200)
                result_workbook = load_workbook(BytesIO(download.data), data_only=True)
                self.assertEqual(result_workbook["小学"]["A2"].value, "待补数据饭")
                self.assertTrue(all(result_workbook["小学"].cell(2, column).value is None for column in range(2, 10)))
                self.assertIn("未知酱料", result_workbook["小学"]["A2"].comment.text)
                result_workbook.close()
                download.close()


if __name__ == "__main__":
    unittest.main()
