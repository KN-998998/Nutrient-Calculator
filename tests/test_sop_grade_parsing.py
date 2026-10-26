import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl import load_workbook

from nutrition_tool.calculator import calculate_recipe
from nutrition_tool.models import NutrientProfile
from nutrition_tool.repository import NutrientRepository
from nutrition_tool.sop import read_sop_workbook


class GradeSpecificSopParsingTests(unittest.TestCase):
    def make_sop(self, folder: str, *, middle_school_grams=200):
        path = Path(folder) / "sample.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "9月A餐配料sop"
        sheet.append([
            "編號/圖片", "套餐", "菜式", "食材/系統名稱", "安士（小學）",
            "人均餐盒實際下單用量（g）（小學）", "安士（中學）",
            "人均餐盒實際下單用量（g）（中學）", "編碼",
        ])
        sheet.append([46266, "雞肉飯", "主菜", "雞肉", 1, 100, 2, middle_school_grams, "CHICKEN"])
        workbook.save(path)
        workbook.close()
        return path

    def test_reads_grade_specific_weights_and_ingredient_code(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self.make_sop(folder)

            recipes, errors = read_sop_workbook(path)
            primary = [recipe for recipe in recipes if recipe.school_stage == "小學"]
            secondary = [recipe for recipe in recipes if recipe.school_stage == "中學"]

        self.assertEqual(errors, [])
        self.assertEqual(primary[0].recipe_name, "雞肉飯")
        self.assertEqual(primary[0].meal_set, "A")
        self.assertEqual(primary[0].ingredients[0].grams, 100)
        self.assertEqual(secondary[0].ingredients[0].grams, 200)
        self.assertEqual(primary[0].ingredients[0].ingredient_code, "CHICKEN")

    def test_includes_school_specific_sop_sheet_only_for_its_named_stage(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self.make_sop(folder)
            workbook = load_workbook(path)
            sheet = workbook.create_sheet("9月D餐配料sop（北角衛理小學）")
            sheet.append([
                "編號/圖片", "套餐", "菜式", "食材/系統名稱", "安士（小學）",
                "人均餐盒實際下單用量（g）（小學）", "安士（中學）",
                "人均餐盒實際下單用量（g）（中學）", "編碼",
            ])
            sheet.append([46266, "北角雞肉飯", "主菜", "雞肉", 1, 100, 2, 200, "CHICKEN"])
            workbook.create_sheet("9月D餐配料sop（測試）")
            workbook.save(path)
            workbook.close()

            recipes, errors = read_sop_workbook(path)

        self.assertEqual(errors, [])
        self.assertEqual(len(recipes), 3)
        special = [recipe for recipe in recipes if recipe.recipe_name == "北角雞肉飯"]
        self.assertEqual(len(special), 1)
        self.assertEqual(special[0].school_stage, "小學")
        self.assertEqual(special[0].ingredients[0].grams, 100)
        self.assertEqual(special[0].source_sheet, "9月D餐配料sop（北角衛理小學）")

    def test_invalid_grade_weight_marks_dish_incomplete_instead_of_dropping_context(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self.make_sop(folder, middle_school_grams="#VALUE!")

            recipes, errors = read_sop_workbook(path)

        secondary = [recipe for recipe in recipes if recipe.school_stage == "中學"]
        self.assertEqual(len(secondary), 1)
        self.assertTrue(secondary[0].parse_warnings)
        self.assertEqual(len(errors), 1)
        self.assertIn("H2", errors[0])

    def test_zero_formula_result_with_blank_grade_amount_is_an_input_error(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self.make_sop(folder, middle_school_grams=0)
            workbook = load_workbook(path)
            workbook.active["G2"] = None
            workbook.save(path)
            workbook.close()

            recipes, errors = read_sop_workbook(path)

        self.assertEqual(len(errors), 1)
        self.assertIn("用量为空", errors[0])
        secondary = [recipe for recipe in recipes if recipe.school_stage == "中學"]
        self.assertEqual(len(secondary), 1)
        profile = NutrientProfile(
            ingredient_code="CHICKEN", ingredient_name="雞肉", energy_kcal=100
        )
        result = calculate_recipe(secondary[0], NutrientRepository([profile]))
        self.assertEqual(result.status, "incomplete")

    def test_count_in_ingredient_name_estimates_only_missing_gram_weights(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "counted.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "9月A餐配料sop"
            sheet.append([
                "編號/圖片", "套餐", "菜式", "食材/系統名稱", "安士（小學）",
                "人均餐盒實際下單用量（g）（小學）", "安士（中學）",
                "人均餐盒實際下單用量（g）（中學）", "編碼",
            ])
            sheet.append([46266, "餐盒", "配菜", "滷水蛋（熟蛋）1個", None, 0, None, 0, "EGG"])
            sheet.append([None, None, "配菜", "麥樂雞4块", None, 33, None, 0, "NUGGET"])
            sheet.append([None, None, "配菜", "蜜糖雞尾腸3個", None, 0, None, 0, "SAUSAGE"])
            workbook.save(path)
            workbook.close()

            recipes, errors = read_sop_workbook(path)

        self.assertEqual(errors, [])
        primary = next(recipe for recipe in recipes if recipe.school_stage == "小學")
        secondary = next(recipe for recipe in recipes if recipe.school_stage == "中學")
        self.assertEqual([item.grams for item in primary.ingredients[:2]], [50, 33])
        self.assertEqual([item.grams for item in secondary.ingredients[:2]], [50, 80])
        self.assertAlmostEqual(primary.ingredients[2].grams, 56.7)
        self.assertAlmostEqual(secondary.ingredients[2].grams, 56.7)
        self.assertTrue(secondary.ingredients[0].weight_is_estimated)
        self.assertTrue(secondary.ingredients[1].weight_is_estimated)
        self.assertFalse(primary.ingredients[1].weight_is_estimated)

    def test_count_only_ingredient_uses_named_component_for_portion_estimate(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "count-only-component.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "9月C餐配料sop"
            sheet.append([
                "編號/圖片", "套餐", "菜式", "食材/系統名稱", "安士（小學）",
                "人均餐盒實際下單用量（g）（小學）", "安士（中學）",
                "人均餐盒實際下單用量（g）（中學）", "編碼",
            ])
            sheet.append([46293, "健康飯卷配甜酸齋", "健康飯卷", "1個", None, 0, None, 0, None])
            workbook.save(path)
            workbook.close()

            recipes, errors = read_sop_workbook(path)

        self.assertEqual(errors, [])
        self.assertEqual(len(recipes), 2)
        for recipe in recipes:
            self.assertEqual(recipe.ingredients[0].ingredient_name, "健康飯卷")
            self.assertEqual(recipe.ingredients[0].grams, 115)
            self.assertTrue(recipe.ingredients[0].weight_is_estimated)
            self.assertIn("vegue.me", recipe.ingredients[0].weight_estimate_note)

    def test_brand_dumpling_count_uses_traceable_brand_pack_weight_for_both_stages(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "brand-dumpling.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "9月E餐配料sop"
            sheet.append([
                "編號/圖片", "套餐", "菜式", "食材/系統名稱", "安士（小學）",
                "人均餐盒實際下單用量（g）（小學）", "安士（中學）",
                "人均餐盒實際下單用量（g）（中學）", "編碼",
            ])
            sheet.append([
                46268, "山東菜肉餃", "山東菜肉餃", "蔥阿伯東北手工水餃",
                "8隻", "#VALUE!", "8隻", "#VALUE!", "CK0226",
            ])
            workbook.save(path)
            workbook.close()

            recipes, errors = read_sop_workbook(path)

        self.assertEqual(errors, [])
        self.assertEqual(len(recipes), 2)
        for recipe in recipes:
            item = recipe.ingredients[0]
            self.assertEqual(item.grams, 144)
            self.assertTrue(item.weight_is_estimated)
            self.assertIn("720g", item.weight_estimate_note)
            self.assertIn("約40", item.weight_estimate_note)

    def test_sop_pack_spec_overrides_generic_weight_for_counted_fish_siu_mai(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "pack-spec.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "9月D餐配料sop"
            sheet.append([
                "編號/圖片", "套餐", "菜式", "食材/系統名稱", "安士（小學）",
                "人均餐盒實際下單用量（g）（小學）", "安士（中學）",
                "人均餐盒實際下單用量（g）（中學）", "編碼", "規格/單位",
            ])
            sheet.append([
                46269, "魚蓉燒賣飯", "魚蓉燒賣4個", "魚蓉燒賣4個", None,
                0, None, 0, "HK0-0644", "2000克/袋，約140粒/袋，8袋/箱",
            ])
            workbook.save(path)
            workbook.close()

            recipes, errors = read_sop_workbook(path)

        self.assertEqual(errors, [])
        self.assertEqual(len(recipes), 2)
        for recipe in recipes:
            item = recipe.ingredients[0]
            self.assertAlmostEqual(item.grams, 2000 / 140 * 4)
            self.assertTrue(item.weight_is_estimated)
            self.assertIn("2000克/袋", item.weight_estimate_note)
            self.assertIn("140粒/袋", item.weight_estimate_note)

    def test_missing_secondary_d_meal_weights_use_documented_portions(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "d-meals.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "9月D餐配料sop"
            sheet.append([
                "編號/圖片", "套餐", "菜式", "食材/系統名稱", "安士（小學）",
                "人均餐盒實際下單用量（g）（小學）", "安士（中學）",
                "人均餐盒實際下單用量（g）（中學）", "編碼",
            ])
            sheet.append([46272, "雞肉漢堡包（蜜糖雞尾腸+車厘茄）", "雞肉漢堡包", "麵包胚", "2個", "#VALUE!", None, 0, None])
            sheet.append([None, None, None, "鸡肉漢堡扒", 3, 85.05, None, 0, None])
            sheet.append([None, None, "車厘茄", "車厘茄", "6个", "#VALUE!", None, 0, "HK0-1353"])
            sheet.append([46276, "熱狗包(法蘭克福腸)（麥樂雞+車厘茄）", "熱狗包(法蘭克福腸)", "麵包胚", "2個", "#VALUE!", None, 0, None])
            sheet.append([None, None, None, "法蘭克福腸", None, 0, None, 0, None])
            sheet.append([None, None, "車厘茄", "車厘茄", "6个", "#VALUE!", None, 0, "HK0-1353"])
            sheet.append([46283, "豬柳漢堡包（蜜糖雞尾腸+車厘茄）", "豬柳漢堡包", "麵包胚", "2個", "#VALUE!", None, 0, None])
            sheet.append([None, None, None, "豬柳漢堡扒", 3, 85.05, None, 0, None])
            sheet.append([None, None, "車厘茄", "車厘茄", "6个", "#VALUE!", None, 0, "HK0-1353"])
            workbook.save(path)
            workbook.close()

            recipes, errors = read_sop_workbook(path)

        self.assertEqual(errors, [])
        secondary = [recipe for recipe in recipes if recipe.school_stage == "中學"]
        self.assertEqual(len(secondary), 3)
        by_name = {recipe.recipe_name: recipe for recipe in secondary}
        for recipe in secondary:
            bun = next(item for item in recipe.ingredients if item.ingredient_name == "麵包胚")
            tomato = next(item for item in recipe.ingredients if item.ingredient_name == "車厘茄")
            self.assertEqual(bun.grams, 103)
            self.assertEqual(tomato.grams, 102)
            self.assertTrue(bun.weight_is_estimated)
            self.assertIn("小學", bun.weight_estimate_note)
            self.assertIn("hanna-one.com", tomato.weight_estimate_note)
        chicken = next(item for item in by_name["雞肉漢堡包（蜜糖雞尾腸+車厘茄）"].ingredients if item.ingredient_name == "鸡肉漢堡扒")
        pork = next(item for item in by_name["豬柳漢堡包（蜜糖雞尾腸+車厘茄）"].ingredients if item.ingredient_name == "豬柳漢堡扒")
        sausage = next(item for item in by_name["熱狗包(法蘭克福腸)（麥樂雞+車厘茄）"].ingredients if item.ingredient_name == "法蘭克福腸")
        self.assertEqual((chicken.grams, pork.grams, sausage.grams), (85.05, 85.05, 100))
        self.assertTrue(all(item.weight_is_estimated for item in (chicken, pork, sausage)))
        self.assertIn("nicolretailer.com", sausage.weight_estimate_note)



if __name__ == "__main__":
    unittest.main()
