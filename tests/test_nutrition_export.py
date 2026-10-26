import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from nutrition_tool.models import IngredientContribution, RecipeResult
from nutrition_tool.nutrition_export import NUTRITION_EXPORT_HEADERS, write_school_nutrition_excel


def _result(name: str, *, status: str = "passed") -> RecipeResult:
    nutrients = {
        "energy_kcal": 123.456,
        "protein_g": 12.345,
        "total_fat_g": 6.789,
        "saturated_fat_g": 1.234,
        "trans_fat_g": 0.012,
        "available_carbohydrate_g": 23.456,
        "sugars_g": 4.567,
        "sodium_mg": 456.7,
        "dietary_fibre_g": 2.0,
    }
    return RecipeResult(
        recipe_name=name,
        meal_set="A",
        service_date="2026-09-01",
        source_sheet="9月A餐配料sop",
        input_weight_g=100,
        edible_weight_g=100,
        estimated_cooked_weight_g=100,
        nutrients_per_serving=nutrients,
        nutrients_per_100g=nutrients,
        macro_energy_kcal=200,
        energy_difference_kcal=0,
        energy_difference_percent=0,
        status=status,
        unmatched_ingredients=[{
            "ingredient_name": "示例缺失食材", "ingredient_code": "X1", "grams": 10,
            "source_row": 5,
        }] if status == "incomplete" else [],
        input_errors=["SOP!F5 克重无法读取"] if status == "incomplete" else [],
    )


class SchoolNutritionExportTests(unittest.TestCase):
    def test_exports_two_sheets_with_only_dish_and_eight_nutrients(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "nutrition.xlsx"
            write_school_nutrition_excel(
                [_result("雞肉飯")], [_result("雞肉飯")], path
            )

            workbook = load_workbook(path, read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames, ["小学", "中学"])
            for sheet in workbook.worksheets[:2]:
                self.assertEqual(tuple(next(sheet.iter_rows(values_only=True))), NUTRITION_EXPORT_HEADERS)
                self.assertEqual(sheet.max_column, 9)
                self.assertEqual(sheet.max_row, 2)
                self.assertEqual(sheet.cell(2, 1).value, "雞肉飯")
                self.assertAlmostEqual(sheet.cell(2, 2).value, 123.456)
            workbook.close()

    def test_incomplete_recipe_exports_blank_nutrients_not_partial_totals(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "nutrition.xlsx"
            write_school_nutrition_excel(
                [_result("待補食材菜式", status="incomplete")], [], path
            )

            workbook = load_workbook(path, read_only=False, data_only=True)
            row = next(workbook["小学"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(row[0], "待補食材菜式")
            self.assertEqual(row[1:], (None,) * 8)
            self.assertEqual(workbook.sheetnames, ["小学", "中学"])
            workbook.close()

    def test_incomplete_main_row_is_highlighted_with_actionable_comment(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "nutrition.xlsx"
            write_school_nutrition_excel(
                [_result("待补数据菜式", status="incomplete")], [], path
            )
            workbook = load_workbook(path)
            row = workbook["小学"]["A2"]
            self.assertEqual(row.fill.fgColor.rgb[-6:], "FCE4D6")
            self.assertIsNotNone(row.comment)
            self.assertIn("示例缺失食材", row.comment.text)
            workbook.close()

    def test_incomplete_comment_identifies_unreported_nutrients(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "nutrition.xlsx"
            result = _result("缺少反式脂肪的菜式", status="incomplete")
            result.unmatched_ingredients = []
            result.input_errors = []
            result.contributions = [IngredientContribution(
                ingredient_code="SEASONING",
                ingredient_name="调味粉",
                source_name="调味粉",
                input_weight_g=5,
                edible_weight_g=5,
                cooked_weight_g=5,
                nutrients={},
                missing_nutrients=("trans_fat_g",),
            )]
            write_school_nutrition_excel([result], [], path)

            workbook = load_workbook(path)
            comment = workbook["小学"]["A2"].comment.text
            self.assertIn("反式脂肪", comment)
            workbook.close()

    def test_review_row_marks_estimate_and_preserves_source(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "nutrition.xlsx"
            result = _result("待复核菜式", status="review")
            result.warnings = ["使用同类食材估算。"]
            result.contributions = [IngredientContribution(
                ingredient_code="ESTIMATE",
                ingredient_name="栗米墨魚餅",
                source_name="栗米墨魚餅",
                input_weight_g=30,
                edible_weight_g=30,
                cooked_weight_g=30,
                nutrients={},
                source_type="reference_estimate",
                source_url="https://example.org/source",
                source_label="同类产品估算",
                requires_review=True,
            )]
            write_school_nutrition_excel([result], [], path)

            workbook = load_workbook(path)
            sheet = workbook["小学"]
            self.assertAlmostEqual(sheet["B2"].value, 123.456)
            self.assertEqual(sheet["A2"].fill.fgColor.rgb[-6:], "FFF2CC")
            self.assertIn("待营养师复核", sheet["A2"].comment.text)
            self.assertIn("栗米墨魚餅", sheet["A2"].comment.text)
            self.assertIn("https://example.org/source", sheet["A2"].comment.text)
            self.assertEqual(sheet.max_column, 9)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
