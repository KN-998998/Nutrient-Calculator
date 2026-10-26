from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from .models import RecipeResult


NUTRIENT_EXPORT_FIELDS = (
    ("energy_kcal", "能量(kcal)"),
    ("protein_g", "蛋白质(g)"),
    ("total_fat_g", "总脂肪(g)"),
    ("saturated_fat_g", "饱和脂肪(g)"),
    ("trans_fat_g", "反式脂肪(g)"),
    ("available_carbohydrate_g", "碳水化合物(g)"),
    ("sugars_g", "糖(g)"),
    ("sodium_mg", "钠(mg)"),
)
NUTRITION_EXPORT_HEADERS = ("菜品名称", *(label for _, label in NUTRIENT_EXPORT_FIELDS))
INCOMPLETE_FILL = PatternFill("solid", fgColor="FCE4D6")
INCOMPLETE_FONT = Font(color="C00000", bold=True)
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")


def _incomplete_details(result: RecipeResult) -> str:
    details = []
    if result.unmatched_ingredients:
        details.append("缺少匹配的食材营养档案：" + "、".join(dict.fromkeys(
            str(item.get("ingredient_name", ""))
            for item in result.unmatched_ingredients
            if item.get("ingredient_name")
        )))
    if result.unmatched_ingredients or result.input_errors:
        details.append("本菜品以下8项营养值均待补齐：" + "、".join(
            label for _, label in NUTRIENT_EXPORT_FIELDS
        ))
    if result.input_errors:
        details.extend(result.input_errors)
    missing_by_ingredient = []
    labels = dict(NUTRIENT_EXPORT_FIELDS)
    for contribution in result.contributions:
        missing = [labels[field] for field in contribution.missing_nutrients if field in labels]
        if missing:
            missing_by_ingredient.append(
                f"{contribution.ingredient_name}（{','.join(missing)}）"
            )
    if missing_by_ingredient:
        details.append("未报告的营养项：" + "、".join(dict.fromkeys(missing_by_ingredient)))
    return "；".join(details)


def _review_details(result: RecipeResult) -> str:
    details = ["此菜品营养值为暂算，待营养师复核后再上架。"]
    for contribution in result.contributions:
        if not (contribution.requires_review or contribution.weight_is_estimated):
            continue
        source = contribution.source_url or contribution.source or "来源待确认"
        details.append(
            f"{contribution.ingredient_name}：{contribution.source_label}；来源：{source}"
        )
        if contribution.weight_is_estimated:
            details.append(
                f"{contribution.ingredient_name}克重估算："
                f"{contribution.weight_estimate_note or '待核对'}"
            )
    details.extend(result.warnings)
    return "\n".join(dict.fromkeys(details))


def _append_results(sheet, results: list[RecipeResult]) -> int:
    sheet.append(NUTRITION_EXPORT_HEADERS)
    count = 0
    for row_number, result in enumerate(results, start=2):
        if result.status == "incomplete":
            values = (None,) * len(NUTRIENT_EXPORT_FIELDS)
        else:
            values = tuple(
                round(result.nutrients_per_serving[field], 3)
                for field, _ in NUTRIENT_EXPORT_FIELDS
            )
        sheet.append((result.recipe_name, *values))
        if result.status == "incomplete":
            for cell in sheet[row_number]:
                cell.fill = INCOMPLETE_FILL
                cell.font = INCOMPLETE_FONT
            missing = "、".join(dict.fromkeys(
                str(item.get("ingredient_name", ""))
                for item in result.unmatched_ingredients
                if item.get("ingredient_name")
            )) or "（无未匹配食材）"
            sheet.cell(row_number, 1).comment = Comment(
                "此菜品营养结果不完整，8项营养值暂不应填写。\n"
                f"待补食材：{missing}\n需补充/核对：{_incomplete_details(result)}",
                "营养素计算工具",
            )
        elif result.status == "review":
            name_cell = sheet.cell(row_number, 1)
            name_cell.fill = REVIEW_FILL
            name_cell.comment = Comment(_review_details(result), "营养素计算工具")
        count += 1
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 42
    for column in "BCDEFGHI":
        sheet.column_dimensions[column].width = 18
    return count


def write_school_nutrition_excel(
    primary_results: Iterable[RecipeResult],
    secondary_results: Iterable[RecipeResult],
    path: str | Path,
) -> dict[str, int]:
    """Write only dish names and the eight Hong Kong nutrition-label values."""
    workbook = Workbook()
    primary = workbook.active
    primary.title = "小学"
    secondary = workbook.create_sheet("中学")
    primary_results = list(primary_results)
    secondary_results = list(secondary_results)
    counts = {
        "primary": _append_results(primary, primary_results),
        "secondary": _append_results(secondary, secondary_results),
    }
    workbook.save(path)
    workbook.close()
    return counts
